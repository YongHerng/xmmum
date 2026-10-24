from io import BytesIO
import pymupdf as fitz
import pytest
from docx import Document
from app.schemas import Answers, CV
from app.engine import evaluate, normalize_skills, matches, forecast, pathways, impact
from app.documents import parse_document
from app.demo import demo_snapshot


def test_aliases_and_word_boundaries():
    assert normalize_skills(["postgresql", "Python programming", "PYTHON", "collaboration"]) == [
        "python",
        "sql",
        "teamwork",
    ]
    assert "machine-learning" not in normalize_skills(["html"])
    assert "sql" not in normalize_skills(["nosql"])


def test_cv_score_is_weighted_and_empty_does_not_score():
    assert evaluate(CV())["score"] == 0
    rich = CV(
        name="Example",
        target_role="data-analyst",
        email="example@example.invalid",
        summary="A student learning data analysis.",
        education="Computer Science",
        projects="Analyzed 2000 readings and created 3 charts for 2 teams.",
        skills=["Python", "SQL", "Excel", "Statistics", "Data analysis"],
    )
    result = evaluate(rich)
    assert result["score"] == round(sum(c["score"] * c["weight"] / 100 for c in result["breakdown"]))
    assert result["score"] > 70
    assert sum(c["weight"] for c in result["breakdown"]) == 100
    assert result["analysis_mode"] == "deterministic"


def test_scenario_does_not_change_personal_fit():
    data = demo_snapshot("student")
    cv = CV(**data["cv"])
    answers = Answers(**data["questionnaire"])
    a = matches(cv, answers, "moderate")
    b = matches(cv, answers, "rapid")
    assert len(a) == 5
    assert [(r["id"], r["score"]) for r in a] == [(r["id"], r["score"]) for r in b]
    assert (
        forecast(cv, "data-analyst", "moderate")["skills"] == forecast(cv, "data-analyst", "rapid")["skills"]
    )
    assert all("constraints" in r["explanation"].lower() for r in a)


def test_pathways_cover_only_missing_skills_and_use_real_catalogue():
    cv = CV(skills=["Python"])
    answers = Answers(weekly_hours=5)
    plans = pathways(cv, answers, "data-analyst")
    assert len(plans) == 2
    for p in plans:
        assert p["hours"] == sum(s["hours"] for s in p["steps"])
        assert p["weeks"] * 5 >= p["hours"]
        assert all(s["url"].startswith("https://") and s["provider"] for s in p["steps"])
        assert all("Python" not in s["covered_skills"] for s in p["steps"])
    assert plans[0]["hours"] <= plans[1]["hours"]


def test_uncovered_resources_are_disclosed():
    plans = pathways(CV(), Answers(), "smart-manufacturing")
    assert "Quality control" in plans[0]["uncovered_skills"]
    with pytest.raises(ValueError):
        pathways(CV(), Answers(), "invented")


def test_outlook_does_not_invent_skills_or_tool_demand():
    assert forecast(CV(), "data-analyst", "rapid")["skills"] == []
    result = forecast(CV(skills=["Python", "Machine learning"]), "data-analyst", "rapid", 10)
    skills = {s["id"]: s for s in result["skills"]}
    assert skills["python"]["direction"] == "Insufficient evidence"
    assert skills["python"]["source"] is None
    assert skills["machine-learning"]["evidence_group"] == "AI and big data"
    assert "weforum.org" in skills["machine-learning"]["source"]["url"]
    assert any("2025–2030 only" in assumption for assumption in result["assumptions"])
    assert all("onetonline.org" in match["occupational_source"]["url"] for match in matches(CV(), Answers()))


@pytest.mark.parametrize("location,currency", [("Malaysia", "MYR"), ("Singapore", "SGD"), ("Penang, Malaysia", "MYR"), ("United States", "USD"), ("Unknown", None)])
def test_pathway_currency_follows_country(location, currency):
    assert pathways(CV(), Answers(location=location), "data-analyst")[0]["currency"] == currency


def test_parse_real_pdf_and_docx_in_memory():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Aisha Example\nEducation\nComputer Science\nProjects\nPython SQL Excel")
    pdf = doc.tobytes()
    doc.close()
    result = parse_document(pdf, "cv.pdf", "application/pdf")
    assert result["cv"]["name"] == "Aisha Example"
    assert "Python" in result["cv"]["skills"]
    word = Document()
    word.add_paragraph("Daniel Example")
    word.add_paragraph("Experience")
    word.add_paragraph("Quality control and troubleshooting")
    stream = BytesIO()
    word.save(stream)
    result = parse_document(
        stream.getvalue(),
        "cv.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "Quality control" in result["cv"]["skills"]
    with pytest.raises(ValueError):
        parse_document(b"fake", "cv.pdf", "application/pdf")
    with pytest.raises(ValueError):
        parse_document(b"%PDF-1.4", "cv.exe", "application/pdf")


def test_demo_is_deterministic_and_explicitly_fictional():
    assert demo_snapshot("student") == demo_snapshot("student")
    assert demo_snapshot("professional")["read_only"]
    assert "Fictional" in demo_snapshot("student")["label"]
    student_demo = demo_snapshot("student")
    assert "skill_mappings" in student_demo["impact"]
    assert "news" in student_demo and len(student_demo["news"]) >= 5
    assert "forum_posts" in student_demo and len(student_demo["forum_posts"]) >= 3
    assert "tracker" in student_demo and student_demo["tracker"]["xp_points"] > 0
