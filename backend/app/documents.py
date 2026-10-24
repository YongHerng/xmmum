import re
from io import BytesIO
from unicodedata import normalize
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document

from .catalog import label
from .document_layout import MAX_PIXELS, document_layout, encode_pdf, evidence, word_preview
from .engine import normalize_skills
from .schemas import CV

MAX_SIZE = 10 * 1024 * 1024


SECTION_HEADINGS = {
    "summary": {
        "summary", "professional summary", "profile", "profile summary", "professional profile",
        "career summary", "career profile", "personal profile", "career objective", "objective",
        "professional objective", "personal statement", "about me",
    },
    "education": {
        "education", "education and qualifications", "qualifications", "academic background",
        "education and training", "academic history", "educational background",
    },
    "experience": {
        "experience", "work experience", "working experience", "professional experience",
        "employment", "employment history", "work history", "internships",
        "projects and extracurricular", "projects and extracurricular activities", "experience and projects",
        "work and project experience", "experience projects and extracurricular activities",
        "extracurricular", "extra curricular", "extra curricular activities", "extracurricular activities",
        "extra curricular achievements", "extracurricular achievements", "co curricular activities",
        "campus involvement", "activities and projects", "volunteer", "volunteering",
        "volunteer experience", "leadership activities",
    },
    "projects": {
        "project experience", "projects", "academic projects", "selected projects", "personal projects", "key projects",
        "achievements", "key achievements", "selected achievements", "accomplishments", "awards",
        "achievements and awards", "awards and achievements", "honors and awards", "honours and awards",
        "honors and achievements", "honours and achievements", "scholarships and awards",
        "awards and scholarships", "recognition and awards",
    },
    "skills": {
        "skills", "technical skills", "key skills", "core skills", "core competencies",
        "key competencies and skills", "competencies and skills", "skills and competencies",
        "areas of expertise", "tools and technologies", "programming languages", "technical skills and languages",
        "skills and languages", "skills and language", "language skills", "languages",
    },
    "membership_licensure": {
        "membership", "memberships", "professional membership", "professional memberships",
        "professional affiliations", "affiliations", "licensure", "license", "licenses", "licence",
        "licences", "professional license", "professional licenses", "professional registration",
        "professional registrations", "membership and licensure", "memberships and licensure",
        "licenses and memberships", "licences and memberships", "professional licenses and memberships",
        "professional affiliations and licenses", "professional affiliations and memberships",
        "professional registrations and memberships",
    },
    "certificates_training": {
        "certification", "certifications", "certificate", "certificates", "professional certifications",
        "courses", "training", "training and development", "professional development",
        "certificate and training", "certificates and training", "certification and training", "certifications and training",
        "training and certificates", "training and certifications", "courses and certifications",
        "courses and training",
    },
    "mixed_credentials": {
        "membership and certification", "membership and certificates",
        "memberships and certificates", "memberships and certifications", "certifications and memberships",
        "professional memberships and certificates", "professional memberships and certifications",
        "certifications and professional memberships",
        "professional affiliations and certifications", "certifications and professional affiliations",
        "license and certification", "licenses and certifications", "licence and certification",
        "licences and certifications", "certifications and licenses", "certifications and licences",
        "license and certifications", "licence and certifications", "certification and licensure",
        "certifications and licensure", "licensure and certifications",
    },
    "other": {
        "honors", "honours", "interests", "hobbies", "references",
        "referees", "publications", "additional information", "activities", "volunteer activities",
    },
}
TITLE_HEADINGS = {"cv", "resume", "résumé", "curriculum vitae", "personal information", "contact", "contact information",
                  "sample cv", "fictional sample cv", "sample resume", "fictional sample resume"}
AWARD_ITEM_PATTERN = re.compile(
    r"\b(?:dean[’']?s?\s+list|scholarship|award|prize|winner|recipient|medal|\d+(?:st|nd|rd|th)\s+place|"
    r"(?:first|second|third)\s+place|top\s+\d+)\b", re.IGNORECASE,
)
DATE_LINE_PATTERN = re.compile(
    r"^(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[ .-]*\s*)?\d{4}"
    r"(?:\s*(?:-|–|—|to)\s*(?:present|current|\d{4}|"
    r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[ .-]*\d{4}))?$",
    re.IGNORECASE,
)
EMAIL_PATTERN = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
PHONE_PATTERN = re.compile(r"(?:\+?\d[\d .()/-]{6,}\d)")
PHONE_LABEL_PATTERN = re.compile(
    r"\b(?:phone|mobile|cell|telephone|tel)[ \t]*(?:number)?[ \t]*[:.]?[ \t]*(\+?\d[\d \t()./-]{6,}\d)",
    re.IGNORECASE,
)


def _clean_line(line: str) -> str:
    line = normalize("NFKC", line)
    line = re.sub(r"[\u200b-\u200f\ufeff]", "", line).replace("\ufffd", "")
    return re.sub(r"^[\s\u2022\u25aa\u25cf\u25e6*\-–—�]+", "", line).strip()


def _normalize_heading(value: str) -> str:
    value = value.casefold().replace("&", " and ").replace("/", " and ")
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def _is_heading_phrase(value: str) -> bool:
    words = re.findall(r"[A-Za-z]+", value)
    if not words or len(words) > 8 or re.search(r"[.!?]", value):
        return False
    title_case = sum(word[0].isupper() for word in words) / len(words)
    return value.isupper() or title_case >= 0.7


def _heading(line: str) -> tuple[str | None, str]:
    cleaned = _clean_line(line)
    match = re.match(r"^([^:]{2,80}):\s*(.*)$", cleaned)
    label_text, remainder = (match.group(1), match.group(2)) if match else (cleaned, "")
    normalized = _normalize_heading(label_text)
    aliases = {_normalize_heading(alias): section for section, values in SECTION_HEADINGS.items() for alias in values}
    if normalized in aliases:
        return aliases[normalized], remainder
    if normalized in {_normalize_heading(alias) for alias in TITLE_HEADINGS}:
        return "title", remainder
    if not _is_heading_phrase(label_text):
        return None, cleaned

    words = set(normalized.split())
    first = next(iter(normalized.split()), "")
    if first in {"summary", "profile", "career", "professional", "personal", "objective"} and words & {"summary", "profile", "objective", "statement"}:
        return "summary", remainder
    if first in {"skill", "skills", "technical", "soft", "key", "core", "language", "programming", "computer", "competency", "competencies", "areas"} and words & {"skill", "skills", "competency", "competencies", "expertise"}:
        return "skills", remainder
    membership_terms = {"membership", "memberships", "licence", "licences", "license", "licenses", "licensure", "affiliation", "affiliations", "registration", "registrations"}
    certificate_terms = {"certification", "certifications", "certificate", "certificates", "training", "course", "courses"}
    if words & membership_terms and words & certificate_terms:
        return "mixed_credentials", remainder
    if first in {"membership", "memberships", "professional", "licence", "licences", "license", "licenses", "licensure", "affiliation", "affiliations", "registration", "registrations"} and words & membership_terms:
        return "membership_licensure", remainder
    if first in {"extra", "extracurricular", "co", "curricular", "campus", "volunteer", "volunteering", "leadership", "activities"} and words & {"extracurricular", "curricular", "campus", "volunteer", "volunteering", "leadership", "activities", "involvement"}:
        return "experience", remainder
    if first in {"project", "projects", "academic", "selected", "personal", "key"} and words & {"project", "projects"}:
        return "projects", remainder
    if first in {"award", "awards", "achievement", "achievements", "accomplishment", "accomplishments", "honor", "honors", "honour", "honours", "recognition"}:
        return "projects", remainder
    if first in {"experience", "employment", "internship", "internships", "work", "professional"} and words & {"experience", "employment", "internship", "internships", "work", "history"}:
        return "experience", remainder
    if first in {"education", "academic"} and words & {"education", "academic", "history", "background"}:
        return "education", remainder
    # Credential names are content, not headings (e.g. "First Aid Certificate").
    if words & certificate_terms and words <= certificate_terms | {
        "and", "professional", "technical", "additional", "relevant", "continuing", "education", "development",
    }:
        return "certificates_training", remainder
    if words & {"interest", "interests", "hobbies", "references", "referees", "publication", "publications"}:
        return "other", remainder
    return None, cleaned


def _wrapped_heading_lines(text: str, source_lines: list[dict] | None = None) -> list[str]:
    """Restore wrapped headings and order same-row side labels before body text."""
    if not source_lines:
        return [_clean_line(line) for line in text.splitlines() if _clean_line(line)]

    records = []
    for index, record in enumerate(source_lines):
        value = _clean_line(str(record.get("text", "")))
        if not value:
            continue
        box = record.get("box", [0, index * 0.02, 1, index * 0.02 + 0.01])
        records.append({"index": index, "page": record.get("page", 0), "text": value, "box": box})

    # Section labels may wrap down the left rail while body text begins at the
    # same height in a right column. Group short uppercase lines by their x position.
    heading_phrases = {_normalize_heading(alias) for values in SECTION_HEADINGS.values() for alias in values}
    heading_words = {word for phrase in heading_phrases for word in phrase.split()}
    candidates = [r for r in records if r["text"].isupper() and len(r["text"]) <= 50
                  and not re.search(r"\d", r["text"])
                  and set(_normalize_heading(r["text"]).split()) <= heading_words]
    candidates.sort(key=lambda r: (r["page"], r["box"][1], r["box"][0]))
    groups: list[list[dict]] = []
    for record in candidates:
        if groups:
            previous = groups[-1][-1]
            vertical_gap = record["box"][1] - previous["box"][3]
            if (record["page"] == previous["page"] and abs(record["box"][0] - previous["box"][0]) <= 0.035
                    and -0.01 <= vertical_gap <= 0.009):
                groups[-1].append(record)
                continue
        groups.append([record])

    consumed: set[int] = set()
    events = []
    for group in groups:
        combined = " ".join(item["text"] for item in group)
        # Combine only a known wrapped heading, never uppercase skill names or
        # adjacent independent headings. Uncombined lines retain their own order.
        if len(group) > 1 and _normalize_heading(combined) not in heading_phrases:
            continue
        section, _ = _heading(combined)
        if section is None and len(group) == 1:
            section, _ = _heading(group[0]["text"])
        if section is None and not _looks_like_unknown_heading(combined):
            continue
        consumed.update(item["index"] for item in group)
        page = group[0]["page"]
        y = max(0.0, min(item["box"][1] for item in group) - 0.006)
        events.append((page, y, 0, min(item["box"][0] for item in group), combined))

    for record in records:
        if record["index"] in consumed:
            continue
        events.append((record["page"], record["box"][1], 1, record["box"][0], record["text"]))
    events.sort(key=lambda item: (item[0], item[1], item[2], item[3]))
    return [event[4] for event in events]


def _looks_like_unknown_heading(line: str) -> bool:
    letters = [char for char in line if char.isalpha()]
    if len(line) > 70 or not letters or re.search(r"[.!?]", line) or re.search(r"\d", line):
        return False
    if line.rstrip().endswith((",", ":", ";")):
        return False
    uppercase = sum(char.isupper() for char in letters) / len(letters)
    words = set(_normalize_heading(line).split())
    heading_cues = {"skills", "competencies", "experience", "education", "projects", "awards", "achievements", "activities", "involvement", "references", "interests", "languages", "certifications", "membership", "license", "licence", "profile", "summary"}
    return uppercase > 0.82 and len(line.split()) <= 8 and bool(words & heading_cues) and not _is_award_item(line)


def _is_award_item(line: str) -> bool:
    return bool(AWARD_ITEM_PATTERN.search(_clean_line(line)))


def _extract_phone(text: str, lines: list[str]) -> str:
    labelled = PHONE_LABEL_PATTERN.search(text)
    candidates = [labelled.group(1)] if labelled else []
    if not candidates:
        for line in lines:
            section = _heading(line)[0]
            if section and section != "title":
                break
            candidates.extend(match.group(0) for match in PHONE_PATTERN.finditer(line))
    for candidate in candidates:
        digits = re.sub(r"\D", "", candidate)
        minimum = 8 if labelled else 9
        if minimum <= len(digits) <= 15:
            return re.sub(r"\s+", " ", candidate).strip()
    return ""


CERTIFICATE_PATTERN = re.compile(
    r"\b(?:certificat(?:e|ion|ions)|certified|licensed|licen[cs]ed|course|training|workshop|"
    r"completed\s+(?:the\s+)?(?:course|training|certification|certificate)|"
    r"credential|professional development|first aid)\b", re.IGNORECASE,
)
MEMBERSHIP_PATTERN = re.compile(
    r"\b(?:member(?:ship)?s?|fellow|affiliation|license|licence|registr(?:ation|ant)|"
    r"professional body|society|institution of|board of|association of|chartered|registered)\b", re.IGNORECASE,
)

def _credential_destination(line: str) -> str | None:
    """Classify only explicit credential or membership evidence in combined sections."""
    certificate = bool(CERTIFICATE_PATTERN.search(line))
    membership = bool(MEMBERSHIP_PATTERN.search(line))
    if certificate == membership:
        return None
    return "certificates_training" if certificate else "membership_licensure"


def _classify_cv(text: str, source_lines: list[dict] | None = None) -> tuple[str, str, str, dict[str, list[str]], list[str] | None]:
    lines = _wrapped_heading_lines(text, source_lines)
    sections = {
        name: [] for name in ("summary", "education", "experience", "projects", "membership_licensure", "certificates_training")
    }
    skill_lines: list[str] = []
    has_skill_section = False
    if not lines:
        return "", "", "", sections, skill_lines

    name_lines = [_clean_line(line) for line in text.splitlines() if _clean_line(line)]
    explicit_name = next((re.sub(r"^(?:full\s+name|name)\s*:", "", line, flags=re.IGNORECASE)
                          for line in name_lines if re.match(r"^(?:full\s+name|name)\s*:", line, re.IGNORECASE)), "")
    email_match = EMAIL_PATTERN.search(text)
    email = email_match.group(0) if email_match else ""
    phone = _extract_phone(text, lines)
    contact_free = []
    for line in name_lines:
        line = EMAIL_PATTERN.sub(" ", line)
        line = PHONE_PATTERN.sub(" ", line)
        line = re.sub(r"https?://\S+|www\.\S+", " ", line, flags=re.IGNORECASE)
        line = re.sub(r"\s*[|·]\s*", " ", line).strip(" ,;")
        if line:
            contact_free.append(line)
    name = explicit_name or next((line for line in contact_free
                                  if _heading(line)[0] not in {*SECTION_HEADINGS.keys(), "title"}
                                  and len(line.split()) >= 2), "")
    if name:
        name = re.sub(r"^(?:full\s+name|name)\s*:", "", name, flags=re.IGNORECASE).strip()

    active = "other"
    previous_destination = "other"
    for original in lines:
        if original == name or re.match(r"^(?:full\s+name|name)\s*:", original, re.IGNORECASE):
            continue
        section, remainder = _heading(original)
        if section == "title":
            active = "other"
            previous_destination = "other"
            continue
        if section:
            active = section
            has_skill_section = has_skill_section or section == "skills"
            if remainder:
                if section == "skills":
                    skill_lines.append(remainder)
                elif section == "mixed_credentials":
                    destination = _credential_destination(remainder)
                    if destination:
                        sections[destination].append(remainder)
                elif section in sections:
                    sections[section].append(remainder)
            previous_destination = section
            continue
        if _looks_like_unknown_heading(original):
            active = "other"
            previous_destination = "other"
            continue
        if active == "other":
            continue
        line = EMAIL_PATTERN.sub("", _clean_line(original)).strip(" |,;")
        if not line:
            continue
        if active == "skills":
            skill_lines.append(line)
            previous_destination = "skills"
            continue

        destination = _credential_destination(line) if active == "mixed_credentials" else active
        if active in {"education", "experience"} and _is_award_item(line):
            destination = "projects"
        elif DATE_LINE_PATTERN.fullmatch(line) and previous_destination == "projects":
            destination = "projects"
        if destination in sections:
            sections[destination].append(line)
        previous_destination = destination or active
    return name, email, phone, sections, skill_lines if has_skill_section else None


def _extract_skills(skill_lines: list[str]) -> list[str]:
    entries: list[str] = []
    category_labels = {
        "skills", "skill", "technical skills", "soft skills", "language skills", "programming languages",
        "computer skills", "computer software and frameworks", "frameworks and technologies",
        "technical skills and languages", "skills and languages", "languages",
    }
    language_labels = {"english", "chinese", "mandarin", "malay", "cantonese", "hakka", "spanish", "french", "arabic", "hindi", "japanese", "korean", "german"}
    prefix_pattern = re.compile(r"^(?:proficient in|skilled in|experienced in using|experienced with|familiar with|knowledge of|expertise in|language)\s+", re.IGNORECASE)
    for line in skill_lines:
        line = _clean_line(line).strip()
        if not line:
            continue
        line = re.sub(r"^[^:]{2,70}:\s*", lambda match: "" if _normalize_heading(match.group(0).rstrip(": ")) in category_labels else match.group(0), line)
        parts = re.split(r"[,;|•·] +|[,;|•·]|\s{3,}", line)
        for part in parts:
            entry = re.sub(r"\s+", " ", part.strip(" -–—·•�")).rstrip(".")
            if not entry or _normalize_heading(entry) in category_labels:
                continue
            language_match = re.match(r"^([A-Za-z][A-Za-z -]{1,24}):\s*(.+)$", entry)
            if language_match and _normalize_heading(language_match.group(1)) in language_labels:
                entry = f"{language_match.group(1).strip()} ({language_match.group(2).strip()})"
            elif prefix_pattern.match(entry):
                entry = prefix_pattern.sub("", entry)
                if " and " in entry.casefold():
                    parts.extend(re.split(r"\s+and\s+", entry, flags=re.IGNORECASE))
                    continue
            if len(entry) > 90 or len(entry.split()) > 10 or re.search(r"[.!?]$", entry):
                continue
            if re.match(r"^(?:fluent|moderate|basic|advanced|intermediate)\s+(?:in |proficiency|spoken|written)", entry, re.IGNORECASE):
                continue
            if re.match(r"^(?:performs|demonstrates|demonstrated|capable of|able to|known for|strong ability to)\b", entry, re.IGNORECASE):
                continue
            entries.append(entry)
    canonical = [label(skill) for skill in normalize_skills(entries)]
    custom = [entry for entry in entries if not normalize_skills([entry])]
    result, seen = [], set()
    for skill in canonical + custom:
        key = skill.casefold()
        if key not in seen:
            seen.add(key)
            result.append(skill)
        if len(result) == 100:
            break
    return result


def parse_document(content: bytes, filename: str, mime: str, include_ai_source: bool = False, ai_primary: bool = False):
    if not content or len(content) > MAX_SIZE:
        raise ValueError("Choose a non-empty PDF, DOCX, JPG or PNG no larger than 10 MB.")
    suffix = filename.lower().rsplit(".", 1)[-1]
    converted = False
    reformatted = False
    pdf = content
    if suffix == "pdf":
        if not content.startswith(b"%PDF-") or mime not in ["application/pdf", "application/octet-stream"]:
            raise ValueError("The file does not contain a valid PDF signature and MIME type.")
    elif suffix == "docx":
        if not content.startswith(b"PK") or mime not in [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"]:
            raise ValueError("The file does not contain a valid DOCX signature and MIME type.")
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 25 * 1024 * 1024 or len(archive.infolist()) > 2000:
                    raise ValueError("The DOCX expands beyond the supported size.")
                if "[Content_Types].xml" not in archive.namelist() or "word/document.xml" not in archive.namelist():
                    raise ValueError("This archive is not a Word document.")
                if any("vbaproject" in n.lower() for n in archive.namelist()):
                    raise ValueError("Macro-enabled documents are not supported.")
            doc = Document(BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
            if len("\n".join(paragraphs)) > 200000:
                raise ValueError("Extracted document text is too long.")
            pdf, reformatted = word_preview(content, paragraphs)
            converted = True
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Unable to read this DOCX. Export it as PDF or enter your details manually.") from exc
    elif suffix in ("png", "jpg", "jpeg"):
        valid = ((suffix == "png" and content.startswith(b"\x89PNG\r\n\x1a\n") and mime in ("image/png", "application/octet-stream"))
                 or (suffix in ("jpg", "jpeg") and content.startswith(b"\xff\xd8\xff") and mime in ("image/jpeg", "application/octet-stream")))
        if not valid:
            raise ValueError("The file does not contain a valid JPG or PNG signature and MIME type.")
        try:
            with fitz.open(stream=content, filetype="png" if suffix == "png" else "jpeg") as image:
                # image_profile(bytes) fails in the locked PyMuPDF release's bindings.
                # Read native image metadata before conversion, without rasterizing it.
                images = image[0].get_image_info()
                if not images or any(info['width'] * info['height'] > MAX_PIXELS for info in images):
                    raise ValueError("Choose an image no larger than 24 megapixels.")
                pdf = image.convert_to_pdf()
                converted = True
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Unable to open this image. Try a different PDF, JPG or PNG.") from exc
    else:
        raise ValueError("Only PDF, DOCX, JPG and PNG are supported.")
    try:
        text, layout = document_layout(pdf)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read this document. Try another file or enter your details manually.") from exc
    if len(text) > 200000:
        raise ValueError("Extracted document text is too long.")
    if not text.strip():
        raise ValueError("No readable text found. Try a clearer scan or enter your details manually.")
    if ai_primary:
        # Reading/OCR only: no legacy section classifier or inferred skill list.
        email = EMAIL_PATTERN.search(text)
        phone = _extract_phone(text,[line.strip() for line in text.splitlines() if line.strip()])
        cv = CV(email=email[0][:200] if email else '', phone=phone[:50])
        return {'cv':cv.model_dump(), 'characters':len(text),
                'document':{'pages':layout['pages'],'fields':[], 'used_ocr':layout['used_ocr'],
                            'reformatted':reformatted,'pdf_base64':encode_pdf(pdf) if converted else None},
                '_ai_source':{'text':text,'lines':layout['lines']},
                'notice':'AI extraction is a draft. Review every field against the original.'}
    name, email, phone, sections, skill_lines = _classify_cv(text, layout["lines"])
    # A CV without a Skills heading can still contain explicit catalogue skills.
    # Keep them reviewable; never infer skills from certificates or memberships.
    inferred_skills = skill_lines is None
    skills = ([label(skill) for skill in normalize_skills([
        "\n".join(sections[key]) for key in ("summary", "education", "experience", "projects")
    ])] if inferred_skills else _extract_skills(skill_lines))
    cv = CV(
        name=name[:100],
        email=email[:200],
        phone=phone[:50],
        summary="\n".join(sections["summary"])[:5000],
        education="\n".join(sections["education"])[:5000],
        experience="\n".join(sections["experience"])[:10000],
        projects="\n".join(sections["projects"])[:10000],
        membership_licensure="\n".join(sections["membership_licensure"])[:5000],
        certificates_training="\n".join(sections["certificates_training"])[:5000],
        skills=skills,
    )
    fields = evidence(cv.model_dump(), layout["lines"])
    if inferred_skills and skills:
        for field in fields:
            if field["key"] == "skills":
                field.update(needs_review=True, reason="Skills found outside a Skills section. Confirm these are skills you can claim.")
    result = {
        "cv": cv.model_dump(),
        "characters": len(text),
        "document": {
            "pages": layout["pages"], "fields": fields,
            "used_ocr": layout["used_ocr"], "reformatted": reformatted,
            "pdf_base64": encode_pdf(pdf) if converted else None,
        },
        "notice": "Review every field. Extraction can miss or misclassify information. Original file discarded.",
    }
    if include_ai_source:
        result["_ai_source"] = {"text": text, "lines": layout["lines"]}
    return result
